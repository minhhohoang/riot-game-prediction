"""Offline preparation and opt-in Pyro NUTS for the completed broad cohort.

No collector, network client, sampler, or optional Bayesian backend runs on
import. Alternative picks never enter the Bernoulli likelihood.
"""
from __future__ import annotations

import hashlib
import inspect
import io
import json
import math
import time
import uuid
from collections import Counter
from dataclasses import asdict, dataclass
from decimal import Decimal
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.linalg import helmert
from scipy.special import expit
from sklearn.metrics import brier_score_loss, log_loss, roc_auc_score

from hlr_model import lolalytics_matchup_scraper as registry
from hlr_model.broad_rank_matchups import atomic_bytes, atomic_json, canonical, require, sha, stamp

KEY = ['puuid', 'match_id']
GROUPS = tuple('ABCDEFGHIJ') + ('Unclassified',)
RANKS = ('BRONZE', 'SILVER', 'GOLD', 'PLATINUM', 'EMERALD', 'DIAMOND')
ROLES = ('top', 'jungle', 'mid', 'adc', 'support')
BINS = ('0', '1–4', '5–9', '10–19', '20–49', '50+')
CONTINUOUS = ('matchup_probability', 'log_total', 'historical_wr_beta11')
SCHEMA = 'broad-grouped-proficiency-nuts-v1'
TIE_RULE = ('Sort positive scores descending. For a tied block use the mean zero-based rank r; '
            'assign floor(9*r/(N_positive-1)), capped at 8, to A–I. A single positive champion goes to A. '
            'All ties stay together; report empty/uneven groups. Exactly zero clipped scores go to J.')
LIMITS = [
    'Recorded prior champion games are season/history-bounded all-role counts, not lifetime experience or mastery points.',
    'External groups use the workbook\'s patch 13.7, mostly Gold, EUW/NA/KR study; transport to this cohort is an assumption.',
    'External gains define groups only and are never added to a predicted probability; J may have a nonzero fitted slope.',
    'Ranks were checked at collection time, not historically verified for patch 16.19.',
    'Full-patch matchup rates make this retrospective benchmark-assisted evaluation, not leakage-free prospective prediction.',
    'Earlier completed held-out games may update saved pregame histories; held-out outcomes never refit parameters.',
    'General predictions set player and player–champion intercepts to zero; they do not marginalize those effects.',
    'Unseen player/champion combinations use population distributions with uncertainty, not invented fitted intercepts.',
    'Observational comparisons do not establish that switching picks or additional practice would cause a win.',
    'Posterior intervals condition on external benchmarks, group assignments and model assumptions; they omit benchmark sampling/transport error.',
    'The likelihood has no additional shared-match residual effect; splits are match-disjoint and evaluation reports distinct matches.',
]


@dataclass(frozen=True)
class Config:
    snapshot: str = '20261027T194148939685Z_3b8d192e'
    metadata_sha256: str = 'b84e8a6a646103541860338668c9de0f69a066f4369f1d7842b3f88067713b00'
    study_file: str = 'hlr_model/data/external/Mastery impact by Champion.xlsx'
    study_sha256: str = 'bc38ee25968522a595690cdcfd6b9ad095c200c238f5c71e6b88a57ee3743d90'
    validation_start: str = '2026-10-02T22:25:06.238+00:00'
    test_start: str = '2026-10-05T01:32:18.312+00:00'
    profile: str = 'full'  # 'smoke' uses the SAME likelihood and NUTS, on a small training subset.
    seed: int = 20261027
    chains: int = 4
    warmup: int = 600
    draws: int = 1200
    target_accept: float = .9
    max_tree_depth: int = 9
    torch_threads: int = 1
    max_rhat: float = 1.01
    min_ess: float = 400

    def __post_init__(self):
        require(self.profile in ('full', 'smoke'), 'Choose full or smoke; no approximate inference fallback.')
        require(pd.Timestamp(self.validation_start).tz is not None and pd.Timestamp(self.test_start).tz is not None,
                'Split cutoffs require explicit timezones.')
        require(pd.Timestamp(self.validation_start) < pd.Timestamp(self.test_start), 'Invalid chronological cutoffs.')
        require(self.chains >= 2 and self.draws >= 4 and self.warmup >= 1 and .5 < self.target_accept < 1,
                'Invalid sampler settings.')

    @property
    def sampling(self):
        if self.profile == 'smoke':
            return dict(chains=2, warmup=25, draws=25, max_tree_depth=4, train_match_limit=24)
        return dict(chains=self.chains, warmup=self.warmup, draws=self.draws,
                    max_tree_depth=self.max_tree_depth, train_match_limit=None)


def load_source(root, config):
    root = Path(root).resolve()
    path = root / 'hlr_model/data/lolalytics/broad_na_16_19_seed_161926_exact_tiers_v1/processed' / config.snapshot
    file = path / 'collection_metadata.json'
    require(sha(file.read_bytes()) == config.metadata_sha256, 'Pinned matchup completion metadata changed.')
    meta = json.loads(file.read_text())
    require(meta.get('collection_complete') is True, 'Select a completed matchup export, not a partial preview.')
    require(meta.get('schema') == 'broad-exact-tier-tables-v1', 'Unknown benchmark export schema.')
    hashes = {}
    for name, digest in meta['output_sha256'].items():
        require(Path(name).name == name, 'Source filename escapes its snapshot.')
        hashes[name] = sha((path / name).read_bytes())
        require(hashes[name] == digest, f'Source hash mismatch: {name}')
    names = ('played_target_games', 'eligible_alternatives', 'matchup_benchmark_lookup')
    frames = {name: pd.read_parquet(path / (name + '.parquet')) for name in names}
    played, alternatives, lookup = [frames[n] for n in names]
    require(not played.duplicated(KEY).any() and played.target_row_id.is_unique, 'Duplicate played identifier.')
    require(not alternatives.duplicated(KEY + ['candidate_champion_id']).any(), 'Duplicate alternative identifier.')
    require('win' not in alternatives and alternatives.candidate_outcome.isna().all(), 'Unplayed alternatives contain outcome labels.')
    require(played.groupby('seed_tier').puuid.nunique().to_dict() == {t:50 for t in RANKS}, 'Expected completed 50-per-tier cohort.')
    require(played.patch.eq('16.19').all() and played.queue.eq(420).all(), 'Wrong observed patch/queue.')
    require(lookup.lookup_key.is_unique and lookup.requested_rank.eq(lookup.tier.str.lower()).all(), 'Duplicate/non-exact-tier lookup.')
    require(lookup.requested_patch.eq('16.19').all() and lookup.requested_region.eq('global').all()
            and lookup.requested_queue_id.eq(420).all(), 'Wrong benchmark filters.')
    for frame in (played, alternatives):
        valid = frame.benchmark_valid_rate
        require(frame.loc[valid, 'benchmark_validation_status'].eq('verified').all(), 'Unverified benchmark marked valid.')
        require(frame.loc[valid, 'benchmark_raw_matchup_winrate_probability'].between(0,1).all(), 'Invalid raw probability.')
        require(np.allclose(frame.loc[valid, 'benchmark_raw_matchup_winrate_pct'] / 100,
                            frame.loc[valid, 'benchmark_raw_matchup_winrate_probability']), 'Percentage/probability disagreement.')
    source_meta = meta['source_snapshot']
    original = root / 'hlr_model/data/runs/broad_na_16_19_seed_161926/processed' / source_meta['snapshot']
    for name, digest in source_meta['files_sha256'].items():
        require(sha((original / name).read_bytes()) == digest, f'Original cohort source changed: {name}')
    require(sha((original / 'collection_metadata.json').read_bytes()) == source_meta['metadata_sha256'], 'Original cohort metadata changed.')
    frames['history'] = pd.read_parquet(original / 'historical_matches.parquet')
    frames['targets'] = pd.read_parquet(original / 'target_matches.parquet')
    identity = dict(snapshot=config.snapshot, metadata_sha256=config.metadata_sha256,
                    output_sha256=hashes, original_cohort=source_meta)
    return dict(path=path, identity=identity, frames=frames)


def champion_names():
    names = {}
    for cid, row in registry.CHAMPIONS.items():
        for name in (row['riot_name'], *registry.DISPLAY_ALIASES.get(cid, ())):
            names.setdefault(name.replace('’', "'").strip().casefold(), set()).add(cid)
    # Reviewed explicit display name in this workbook, corresponding to Riot Nunu (20).
    names['nunu & willump'] = {20}
    return names


def assign_groups(rows):
    """Use the full usable external table, never just champions seen in our cohort."""
    out = rows.copy()
    out['mastery_group'] = 'Unclassified'
    eligible = out.study_usable & out.champion_id.notna() & out.grouping_score.notna()
    out.loc[eligible & out.grouping_score.eq(0), 'mastery_group'] = 'J'
    positive = out.loc[eligible & out.grouping_score.gt(0)]
    higher = 0
    for score in sorted(positive.grouping_score.unique(), reverse=True):
        ids = positive.index[positive.grouping_score.eq(score)]
        midpoint = higher + (len(ids) - 1) / 2
        index = min(8, int(math.floor(9 * midpoint / max(1, len(positive) - 1))))
        out.loc[ids, 'mastery_group'] = 'ABCDEFGHI'[index]
        higher += len(ids)
    return out


def load_groups(root, config):
    import openpyxl
    file = Path(root) / config.study_file
    base = pd.DataFrame([dict(champion_id=cid, canonical_champion=v['riot_name']) for cid,v in registry.CHAMPIONS.items()])
    if not file.is_file():
        base['mastery_group'] = 'Unclassified'
        base['classification_reason'] = 'external_study_absent'
        for name in ('raw_gain_pp','grouping_score','study_medium_probability','study_high_probability'):
            base[name] = np.nan
        return base, pd.DataFrame(), dict(available=False, modeling_ready=False, expected_path=str(file), tie_rule=TIE_RULE)
    require(sha(file.read_bytes()) == config.study_sha256, 'Workbook differs from the pinned inspected source; review and repin explicitly.')
    book = openpyxl.load_workbook(file, data_only=True, read_only=False)
    require({'Data Intro','Easiest to Learn','Best to Master'} <= set(book.sheetnames), 'Expected the full inspected workbook.')
    sheet = book['Best to Master']
    headers = {c.value:c.column for c in sheet[1] if c.value is not None}
    expected = ('Most Common Lane','Champion Name','Medium Mastery Win Rate','High Mastery Win Rate','High Mastery Ratio')
    require(set(expected) <= set(headers), 'Mastery worksheet columns changed.')
    intro = [str(row[0]) for row in book['Data Intro'].iter_rows(min_col=3,max_col=3,values_only=True) if row[0] is not None]
    links = [c.hyperlink.target for row in book['Data Intro'] for c in row if c.hyperlink]
    source_url = next((u for u in links if 'jackjgaming.substack.com/' in u), None)
    require(source_url is not None, 'Workbook study provenance is absent.')
    aliases = champion_names()
    records = []
    for r in range(2, sheet.max_row + 1):
        name = sheet.cell(r, headers['Champion Name']).value
        if name is None:
            continue
        cells = [sheet.cell(r, headers[k]) for k in ('Medium Mastery Win Rate','High Mastery Win Rate')]
        candidates = aliases.get(str(name).replace('’', "'").strip().casefold(), set())
        cid = next(iter(candidates)) if len(candidates) == 1 else None
        vals, reason = [], None
        for cell in cells:
            value = cell.value
            if isinstance(value, str) and 'low data' in value.lower():
                vals.append(None); reason = 'low_data'
            elif type(value) not in (int,float) or not math.isfinite(value):
                vals.append(None); reason = reason or 'missing_or_invalid_rate'
            else:
                # Excel % formatting establishes that underlying values are probabilities.
                require('%' in cell.number_format and 0 <= value <= 1, 'Ambiguous mastery rate units or invalid probability.')
                vals.append(Decimal(str(value)))
        medium, high = vals
        gain = (high - medium) * 100 if high is not None and medium is not None else None
        usable = gain is not None and cid is not None
        records.append(dict(champion_id=cid, study_champion=str(name), study_row=r,
            study_lane=sheet.cell(r, headers['Most Common Lane']).value,
            study_medium_original=str(cells[0].value), study_high_original=str(cells[1].value),
            medium_excel_format=cells[0].number_format, high_excel_format=cells[1].number_format,
            study_medium_probability=float(medium) if medium is not None else np.nan,
            study_high_probability=float(high) if high is not None else np.nan,
            study_high_mastery_ratio_original=str(sheet.cell(r,headers['High Mastery Ratio']).value),
            raw_gain_pp=float(gain) if gain is not None else np.nan,
            grouping_score=float(max(Decimal(0),gain)) if gain is not None else np.nan,
            study_usable=usable, classification_reason=reason or ('usable' if cid is not None else 'unresolved_or_ambiguous_name'),
            study_source_url=source_url, study_sheet=sheet.title, study_file_sha256=config.study_sha256))
    book.close()
    raw = pd.DataFrame(records)
    require(not raw.loc[raw.champion_id.notna(), 'champion_id'].duplicated().any(), 'Duplicate external champion ID; do not merge ambiguous rows.')
    unresolved = raw.loc[raw.champion_id.isna(), 'study_champion'].tolist()
    require(not unresolved, f'Unresolved study names require explicit reviewed ID mappings: {unresolved}')
    classified = assign_groups(raw)
    lookup = base.merge(classified, on='champion_id', how='outer', validate='one_to_one')
    lookup['mastery_group'] = lookup.mastery_group.fillna('Unclassified')
    lookup['classification_reason'] = lookup.classification_reason.fillna('absent_from_external_study')
    metadata = dict(available=True, modeling_ready=bool(raw.study_usable.any()), file=str(file.resolve()),
        sha256=config.study_sha256, worksheets=['Data Intro','Easiest to Learn','Best to Master'],
        source_url=source_url, intro=intro, tie_rule=TIE_RULE,
        units='Excel cells store probabilities; raw_gain_pp = 100*(high-medium); grouping_score=max(0,raw_gain_pp)',
        source_rows=len(raw), usable_rows=int(raw.study_usable.sum()), positive_rows=int(raw.grouping_score.gt(0).sum()),
        full_study_group_counts=classified.mastery_group.value_counts().reindex(GROUPS,fill_value=0).to_dict(),
        registry_group_counts=lookup.mastery_group.value_counts().reindex(GROUPS,fill_value=0).to_dict(),
        external_gain_in_prediction=False)
    return lookup, classified, metadata


def audit_histories(played, alternatives, history, targets):
    """Reconstruct the existing expanding histories and exact same-role candidate sets.

    Held-out outcomes may enter *later* completed histories; no fitting occurs.
    Source alternatives lack all-role win counts, so reconstruct those using the
    same audited chronology rather than borrowing the played champion's features.
    """
    p = played.copy();a = alternatives.copy()
    require(not p.duplicated(KEY).any() and not a.duplicated(KEY+['candidate_champion_id']).any(), 'Duplicate source rows.')
    require(not history.duplicated(KEY).any() and not targets.duplicated(KEY).any(), 'Duplicate history/target records.')
    require(not set(map(tuple,history[KEY].values)) & set(map(tuple,targets[KEY].values)), 'History and target overlap.')
    anchors = targets.groupby('puuid').timestamp.min()
    require(history.end_timestamp.lt(history.puuid.map(anchors)).all(), 'Initial history is not strictly before first target.')
    originals = p.set_index(KEY)
    require(set(originals.index) == set(map(tuple,targets[KEY].values)), 'Observed/target identifier mismatch.')
    mapped = {key:rows for key,rows in a.groupby(KEY,sort=False)}
    hist = {key:rows for key,rows in history.groupby('puuid',sort=False)}
    p_same, a_wins, a_latest, a_total, a_wr = {}, {}, {}, {}, {}
    for player, group in targets.groupby('puuid',sort=True):
        pool = pd.concat([hist.get(player,history.iloc[:0]),group],ignore_index=True).sort_values(['end_timestamp','match_id'])
        available = list(pool.itertuples(index=False))
        counts, wins, rc, latest = Counter(), Counter(), Counter(), {}
        cursor = total = 0
        for t in group.sort_values(['timestamp','match_id']).itertuples(index=False):
            while cursor < len(available) and available[cursor].end_timestamp < t.timestamp:
                prior = available[cursor]
                require(prior.timestamp < t.timestamp and prior.win in (0,1), 'Invalid prior record chronology/outcome.')
                counts[prior.champion_id] += 1;wins[prior.champion_id] += prior.win
                rc[(prior.lane,prior.champion_id)] += 1
                latest[prior.champion_id] = prior.end_timestamp
                total += 1;cursor += 1
            key = (player,t.match_id);saved = originals.loc[key]
            main = min(counts,key=lambda c:(-counts[c],c)) if counts else None
            require(int(saved.prior_champion_games) == counts[t.champion_id] and int(saved.prior_total_games) == total,
                    'Recorded proficiency counts differ from preserved history.')
            require((pd.isna(saved.prior_main_champion_id) and main is None) or saved.prior_main_champion_id == main,
                    'Existing pregame all-role main differs.')
            require(np.isclose(saved.prior_champion_wr_beta11,(wins[t.champion_id]+1)/(counts[t.champion_id]+2)), 'Historical WR feature differs.')
            p_same[key] = rc[(t.lane,t.champion_id)]
            candidates = mapped.get(key,a.iloc[:0])
            known = not t.matchup_missing and t.lane in ROLES and pd.notna(t.opponent_champion_id) and t.champion_id != t.opponent_champion_id
            expected = {cid for (lane,cid),n in rc.items() if lane==t.lane and n>0} - {t.champion_id,t.opponent_champion_id} if known else set()
            require(set(candidates.candidate_champion_id) == expected, 'Saved alternatives differ from strictly prior same-role eligibility.')
            for idx,row in candidates.iterrows():
                cid = row.candidate_champion_id
                require(row.prior_all_role_games == counts[cid] and row.prior_same_role_games == rc[(t.lane,cid)], 'Alternative count mismatch.')
                require(row.latest_same_role_end < t.timestamp and latest[cid] < t.timestamp, 'Alternative feature chronology failed.')
                a_wins[idx]=wins[cid];a_latest[idx]=latest[cid];a_total[idx]=total
                a_wr[idx]=(wins[cid]+1)/(counts[cid]+2)
    p['same_role_prior_games'] = [p_same[k] for k in p[KEY].itertuples(index=False,name=None)]
    a['all_role_prior_wins'] = pd.Series(a_wins)
    a['latest_all_role_prior_end_timestamp'] = pd.Series(a_latest)
    a['prior_total_games'] = pd.Series(a_total)
    a['historical_wr_beta11'] = pd.Series(a_wr)
    require(a.latest_all_role_prior_end_timestamp.lt(a.timestamp).all(), 'Missing/invalid alternative chronology.')
    return p,a


def split_assignments(played, config):
    cut1 = pd.Timestamp(config.validation_start).value // 10**6
    cut2 = pd.Timestamp(config.test_start).value // 10**6
    require(played.groupby('match_id').timestamp.nunique().le(1).all(), 'Shared match has inconsistent kickoff timestamps.')
    out = played[KEY+['target_row_id','timestamp','end_timestamp']].copy()
    out['partition'] = np.select([out.timestamp.lt(cut1),out.timestamp.lt(cut2)],['train','validation'],default='test')
    out['split'] = out.partition
    for label, cutoff in [('train',cut1),('validation',cut2)]:
        pending = out.loc[out.partition.eq(label)&out.end_timestamp.ge(cutoff),'match_id'].unique()
        out.loc[out.partition.eq(label)&out.match_id.isin(pending),'split'] = 'purged_'+label+'_boundary'
    require(out.groupby('match_id').split.nunique().eq(1).all(), 'A shared match crosses validation partitions.')
    require(out.loc[out.split.eq('train'),'end_timestamp'].lt(cut1).all(), 'Unfinished training outcomes at freeze.')
    require(out.loc[out.split.eq('validation'),'end_timestamp'].lt(cut2).all(), 'Unfinished validation outcomes at test boundary.')
    return out


def prepare(bundle, groups, config):
    f = bundle['frames']
    p,a = audit_histories(f['played_target_games'],f['eligible_alternatives'],f['history'],f['targets'])
    split = split_assignments(p,config)
    lookup = groups[['champion_id','mastery_group','raw_gain_pp','grouping_score','classification_reason']]
    require(lookup.champion_id.is_unique, 'Champion grouping join would multiply rows.')
    p = p.merge(lookup,on='champion_id',how='left',validate='many_to_one')
    a = a.merge(lookup.rename(columns={'champion_id':'candidate_champion_id'}),on='candidate_champion_id',how='left',validate='many_to_one')
    p = p.merge(split[KEY+['partition','split']],on=KEY,validate='one_to_one')
    a = a.merge(split[KEY+['partition','split']],on=KEY,how='left',validate='many_to_one')
    for frame, kind in ((p,'played'),(a,'alternative')):
        require(frame.mastery_group.notna().all(), 'A data champion is absent from the stable-ID registry; review explicitly.')
        frame['row_kind']=kind
        frame['model_rank'] = frame.seed_tier if kind=='played' else frame.tier
        frame['model_champion_id'] = frame.champion_id if kind=='played' else frame.candidate_champion_id
        frame['n_recorded'] = frame.prior_champion_games if kind=='played' else frame.prior_all_role_games
        if kind=='played':frame['historical_wr_beta11']=frame.prior_champion_wr_beta11
        frame['log_total'] = np.log1p(frame.prior_total_games)
        frame['matchup_probability'] = frame.benchmark_raw_matchup_winrate_probability.where(frame.benchmark_valid_rate)
        numeric = frame[['n_recorded',*CONTINUOUS]].apply(pd.to_numeric,errors='coerce')
        reasons=[]
        for idx,row in frame.iterrows():
            why=[]
            if not row.benchmark_valid_rate or not pd.notna(row.matchup_probability):why.append('missing_or_unverified_matchup')
            if row.lane not in ROLES or row.model_rank not in RANKS:why.append('unknown_rank_or_role')
            if not np.isfinite(numeric.loc[idx].to_numpy()).all() or row.n_recorded<0:why.append('invalid_pregame_feature')
            reasons.append('|'.join(why) or 'none')
        frame['data_exclusion_reasons']=reasons
        frame['usable_data']=frame.data_exclusion_reasons.eq('none')
        frame['included_in_primary_model'] = frame.usable_data & ~frame.split.str.startswith('purged') if kind=='played' else False
        frame['experience_bin'] = pd.cut(frame.n_recorded,[-.5,.5,4.5,9.5,19.5,49.5,np.inf],labels=BINS)
    require(len(p)==len(f['played_target_games']) and len(a)==len(f['eligible_alternatives']), 'Preparation lost source rows.')
    require('win' not in a and a.candidate_outcome.isna().all(), 'Alternative outcome contamination.')
    return dict(played=p,alternatives=a,split=split)


def coverage_tables(prepared, groups):
    p,a = prepared['played'],prepared['alternatives']
    rank_role = p.groupby(['seed_tier','lane'],dropna=False).agg(players=('puuid','nunique'),games=('match_id','size'),champions=('champion_id','nunique')).reset_index()
    support = p.groupby('mastery_group').agg(champions=('champion_id','nunique'),players=('puuid','nunique'),games=('match_id','size'),
        median_experience=('n_recorded','median'),max_experience=('n_recorded','max')).reindex(GROUPS).reset_index()
    support[['champions','players','games']]=support[['champions','players','games']].fillna(0).astype(int)
    bins = p.groupby(['mastery_group','experience_bin'],observed=False).agg(games=('match_id','size'),players=('puuid','nunique'),champions=('champion_id','nunique')).reset_index()
    bins = bins.set_index(['mastery_group','experience_bin']).reindex(pd.MultiIndex.from_product([GROUPS,BINS],names=['mastery_group','experience_bin']),fill_value=0).reset_index()
    history = p[['puuid','seed_tier','initial_history_games']].drop_duplicates()
    require(history.puuid.is_unique,'Inconsistent initial history depths.')
    bench=[];exclusions=[]
    for f in (p,a):
        bench.append(f.groupby(['row_kind','model_rank','lane','benchmark_status'],dropna=False).agg(rows=('match_id','size'),players=('puuid','nunique'),sparse_rows=('benchmark_sparse','sum')).reset_index())
        exclusions.append(f.groupby(['row_kind','data_exclusion_reasons'],dropna=False).agg(rows=('match_id','size'),players=('puuid','nunique'),matches=('match_id','nunique')).reset_index())
    unclassified=[]
    for f,kind in ((p,'played'),(a,'alternative')):
        cid='champion_id' if kind=='played' else 'candidate_champion_id'
        unclassified.append(f.loc[f.mastery_group.eq('Unclassified')].groupby(cid).agg(rows=('match_id','size'),players=('puuid','nunique')).reset_index().rename(columns={cid:'champion_id','rows':kind+'_rows','players':kind+'_players'}))
    unknown = groups.loc[groups.mastery_group.eq('Unclassified'),['champion_id','canonical_champion','classification_reason']]
    for f in unclassified:unknown=unknown.merge(f,on='champion_id',how='left',validate='one_to_one')
    unknown=unknown.loc[unknown[['played_rows','alternative_rows']].fillna(0).sum(axis=1).gt(0)].fillna({'played_rows':0,'played_players':0,'alternative_rows':0,'alternative_players':0})
    splits = p.groupby('split').agg(rows=('match_id','size'),players=('puuid','nunique'),matches=('match_id','nunique'),usable_rows=('usable_data','sum'),included_rows=('included_in_primary_model','sum')).reset_index()
    return dict(rank_role_support=rank_role,group_support=support,experience_bins=bins,history_depth=history,
                experience_distribution=p.groupby('mastery_group').n_recorded.describe(percentiles=[.1,.25,.5,.75,.9]).reset_index(),
                matchup_coverage=pd.concat(bench,ignore_index=True),data_exclusions=pd.concat(exclusions,ignore_index=True),
                unclassified_champions=unknown,split_support=splits)


def save_eda_figures(out, prepared, tables):
    import matplotlib.pyplot as plt
    out=Path(out);out.mkdir(parents=True,exist_ok=True)
    paths=[]
    def save(fig,name):
        fig.tight_layout()
        for suffix in ('png','svg'):fig.savefig(out/(name+'.'+suffix),dpi=140,bbox_inches='tight')
        paths.append(out/(name+'.png'));plt.close(fig)
    p=prepared['played']
    fig,axes=plt.subplots(1,2,figsize=(12,4))
    for ax,value,title in zip(axes,['games','players'],['Observed target games','Distinct players per rank/role']):
        f=tables['rank_role_support'].pivot(index='seed_tier',columns='lane',values=value).reindex(RANKS)
        f.plot.bar(ax=ax);ax.set(title=title,xlabel='Collection-time rank');ax.tick_params(axis='x',rotation=30)
        ax.legend(fontsize=7)
    save(fig,'rank_role_support')
    fig,axes=plt.subplots(1,2,figsize=(12,4))
    f=tables['group_support'].set_index('mastery_group')
    f.games.fillna(0).plot.bar(ax=axes[0],title='Observed games by external group')
    f.players.fillna(0).plot.bar(ax=axes[1],title='Distinct players by group (players can span groups)')
    save(fig,'group_support')
    fig,axes=plt.subplots(1,2,figsize=(12,4))
    active=[g for g in GROUPS if p.mastery_group.eq(g).any()]
    axes[0].boxplot([np.log1p(p.loc[p.mastery_group.eq(g),'n_recorded']) for g in active],showfliers=False)
    axes[0].set_xticks(range(1,len(active)+1),active)
    axes[0].set(title='Recorded all-role experience by group',ylabel='log(1 + prior champion games)')
    axes[0].tick_params(axis='x',rotation=35)
    axes[1].hist(tables['history_depth'].initial_history_games,bins=np.arange(0,181,10),color='#4878a8')
    axes[1].set(title='Initial history depth — one row per player',xlabel='Initial historical games (not lifetime)',ylabel='Players')
    save(fig,'experience_and_history')
    fig,axes=plt.subplots(1,2,figsize=(12,5))
    for ax,value in zip(axes,['games','players']):
        matrix=tables['experience_bins'].pivot(index='mastery_group',columns='experience_bin',values=value).reindex(index=GROUPS,columns=BINS).fillna(0)
        image=ax.imshow(np.log1p(matrix.to_numpy(dtype=float)),aspect='auto',cmap='Blues')
        for i in range(len(matrix)):
            for j in range(len(BINS)):ax.text(j,i,str(int(matrix.iloc[i,j])),ha='center',va='center',fontsize=7)
        ax.set(xticks=range(len(BINS)),xticklabels=BINS,yticks=range(len(GROUPS)),yticklabels=GROUPS,title='Experience support: '+value)
        fig.colorbar(image,ax=ax,label='log(1 + count)')
    save(fig,'experience_bin_support')
    fig,axes=plt.subplots(1,2,figsize=(12,4))
    for ax,frame,label in zip(axes,[prepared['played'],prepared['alternatives']],['Played','Alternative']):
        labels=['Missing benchmark','Reported n < 100','Reported n ≥ 100']
        values=[int((~frame.benchmark_valid_rate).sum()),int((frame.benchmark_valid_rate&frame.benchmark_sparse).sum()),int((frame.benchmark_valid_rate&~frame.benchmark_sparse).sum())]
        ax.bar(labels,values,color=['#999999','#ddaa44','#4878a8']);ax.set(title=label+' matchup coverage',ylabel='Rows')
        ax.tick_params(axis='x',rotation=20)
    save(fig,'matchup_coverage')
    return paths


def prepare_outputs(root, bundle, groups, study_rows, study_metadata, prepared, tables, config):
    root=Path(root)/'hlr_model/data/analyses/broad_rank_mastery_v1'
    identity=dict(schema=SCHEMA,source=bundle['identity'],study_sha256=study_metadata.get('sha256'),
                  config=asdict(config),helper_sha256=sha(Path(__file__).read_bytes()))
    identity=canonical(identity)
    for manifest in sorted(root.glob('*/preparation_manifest.json'),reverse=True):
        meta=json.loads(manifest.read_text())
        if meta['identity']==identity:
            require(all(sha((manifest.parent/n).read_bytes())==v for n,v in meta['output_sha256'].items()), 'Prepared analysis artifacts changed; preserve them and use a new version.')
            return manifest.parent
    out=root/(pd.Timestamp.now(tz='UTC').strftime('%Y%m%dT%H%M%S%fZ')+'_'+uuid.uuid4().hex[:8])
    out.mkdir(parents=True,exist_ok=False)
    for name,frame in dict(champion_grouping_lookup=groups,external_study_rows=study_rows,
                           prepared_played=prepared['played'],prepared_alternatives=prepared['alternatives'],
                           split_assignments=prepared['split'],**tables).items():
        frame.to_parquet(out/(name+'.parquet'),index=False)
        if name not in ('prepared_played','prepared_alternatives'):frame.to_csv(out/(name+'.csv'),index=False)
    atomic_json(out/'grouping_metadata.json',study_metadata)
    atomic_json(out/'model_configuration.json',dict(config=asdict(config),sampling=config.sampling,
        specifications=model_specifications(),limits=LIMITS,default_RUN_MODELING=False,
        execution_state_artifacts='fits/*/fit_status.json and modeling_complete.json'))
    atomic_json(out/'source_metadata.json',bundle['identity'])
    save_eda_figures(out/'figures',prepared,tables)
    hashes={str(p.relative_to(out)):sha(p.read_bytes()) for p in out.rglob('*') if p.is_file()}
    atomic_json(out/'preparation_manifest.json',dict(identity=identity,output_sha256=hashes,created_at=stamp(),status='prepared_only'))
    return out


def model_specifications():
    return dict(likelihood='Bernoulli(logit^-1(eta)), played training rows only',
        eta='alpha + standardized [raw matchup probability, log1p(prior total games), all-role prior champion WR Beta(1,1)] + rank + role + champion baseline + player intercept + player–champion intercept + beta_experience[group]*n/(n+k)',
        shared='one beta_experience for all champions',grouped='shared beta_experience plus partially pooled A–J/Unclassified deviations; no ordering constraint',
        k='one shared k across groups per model; log(k) ~ Normal(log(30),1), learned only from training outcomes',
        priors=dict(alpha_and_continuous='Normal(0,2)',rank_role='Normal(0,1)',experience='Normal(0,1)',
                    log_player_and_champion_scale='Normal(log(.3),.7)',log_player_champion_scale='Normal(log(.15),.5)',
                    log_group_contrast_scale='Normal(log(.3),.7)',noncentered_latents='Normal(0,1)'),
        identification='Bronze/top reference levels; zero-sum orthonormal champion baselines and group slope contrasts; proper zero-mean player and player–champion distributions',
        inference='Pyro 1.9.2 NUTS, exact Bernoulli likelihood; no VI/Laplace fallback',
        general_prediction='player and player–champion effects set to zero, not marginalized',
        unclassified='one partially pooled group slope with posterior uncertainty; no independently estimated champion-specific curves',
        adjustment='raw probability is an estimated-slope covariate, never an offset; tier-centered percentage is not passed to logit',
        changes_from_master='Extend the existing nonlinear C model with rank controls, pooled champion baselines and optional group slopes; retain all-role n, historical WR, log_total, priors and cutoff dates; use NUTS for both primary specifications instead of mean-field VI')


class Design:
    """Training-only preprocessing and identity maps; protocol-fixed rank/role levels."""
    def __init__(self, train):
        require(len(train)>0 and train.row_kind.eq('played').all() and 'win' in train,
                'Only observed played training rows can construct the labeled design.')
        require(train.split.eq('train').all() and train.usable_data.all(), 'Design received held-out or unusable rows.')
        require(train.win.isin([0,1]).all(), 'Invalid observed labels.')
        self.mean=train[list(CONTINUOUS)].mean().to_numpy(dtype=float)
        self.scale=train[list(CONTINUOUS)].std(ddof=0).to_numpy(dtype=float)
        self.scale=np.where(self.scale>0,self.scale,1.)
        self.players=sorted(train.puuid.unique())
        self.champions=sorted(int(c) for c in train.model_champion_id.unique())
        self.pcs=sorted(set(zip(train.puuid,train.model_champion_id.astype(int))))
        self.pi={p:i for i,p in enumerate(self.players)}
        self.ci={c:i for i,c in enumerate(self.champions)}
        self.pci={c:i for i,c in enumerate(self.pcs)}
        self.champion_basis=helmert(len(self.champions),full=False).T if len(self.champions)>1 else np.zeros((1,0))
        self.group_basis=helmert(len(GROUPS),full=False).T

    def matrix(self, frame):
        require(frame.model_rank.isin(RANKS).all() and frame.lane.isin(ROLES).all(), 'Unknown rank/role in prediction design.')
        require(frame.mastery_group.isin(GROUPS).all(), 'Unknown mastery group.')
        x=(frame[list(CONTINUOUS)].to_numpy(dtype=float)-self.mean)/self.scale
        require(np.isfinite(x).all() and frame.n_recorded.ge(0).all(), 'Invalid prediction features; no benchmark imputation.')
        return dict(X=x,n=frame.n_recorded.to_numpy(dtype=float),
            rank=np.eye(len(RANKS))[frame.model_rank.map(dict(zip(RANKS,range(len(RANKS))))).to_numpy()][:,1:],
            role=np.eye(len(ROLES))[frame.lane.map(dict(zip(ROLES,range(len(ROLES))))).to_numpy()][:,1:],
            group=frame.mastery_group.map(dict(zip(GROUPS,range(len(GROUPS))))).to_numpy(dtype=int),
            player=np.array([self.pi.get(p,-1) for p in frame.puuid]),
            pc=np.array([self.pci.get((p,int(c)),-1) for p,c in zip(frame.puuid,frame.model_champion_id)]),
            champion=np.array([self.ci.get(int(c),-1) for c in frame.model_champion_id]))

    def metadata(self):
        return dict(continuous=list(CONTINUOUS),mean=self.mean.tolist(),scale=self.scale.tolist(),
                    players=self.players,champions=self.champions,player_champions=[list(k) for k in self.pcs],
                    ranks=RANKS,roles=ROLES,groups=GROUPS,preprocessing='training only; rank/role levels fixed by protocol')


def pyro_model(data, dimensions, grouped=False):
    """Exact generative model; imported backend is required only when explicitly used."""
    import torch
    import pyro
    import pyro.distributions as dist
    def normal(name, loc=0., scale=1., size=None):
        shape=() if size is None else (size,)
        distribution=dist.Normal(torch.full(shape,float(loc),dtype=torch.float64),
                                 torch.full(shape,float(scale),dtype=torch.float64))
        return pyro.sample(name,distribution.to_event(1) if size is not None else distribution)
    alpha=normal('alpha',scale=2.)
    beta=normal('beta',scale=2.,size=len(CONTINUOUS))
    rank=normal('beta_rank',size=len(RANKS)-1)
    role=normal('beta_role',size=len(ROLES)-1)
    beta_exp=normal('beta_exp')
    k=normal('log_k',np.log(30),1.).exp()
    sigma_player=normal('log_sigma_player',np.log(.3),.7).exp()
    sigma_pc=normal('log_sigma_pc',np.log(.15),.5).exp()
    sigma_champion=normal('log_sigma_champion',np.log(.3),.7).exp()
    player=normal('player_z',size=dimensions['players'])*sigma_player
    pc=normal('pc_z',size=dimensions['pcs'])*sigma_pc
    if dimensions['champions']>1:
        champion=data['champion_basis']@normal('champion_free',size=dimensions['champions']-1)*sigma_champion
    else:
        champion=torch.zeros(1,dtype=torch.float64)
    slope=beta_exp.expand(len(GROUPS))
    if grouped:
        scale=normal('log_sigma_group',np.log(.3),.7).exp()
        slope=slope+scale*(data['group_basis']@normal('group_free',size=len(GROUPS)-1))
    eta=(alpha+data['X']@beta+data['rank']@rank+data['role']@role+champion[data['champion']]
         +player[data['player']]+pc[data['pc']]+slope[data['group']]*data['n']/(data['n']+k))
    pyro.sample('outcomes',dist.Bernoulli(logits=eta).to_event(1),obs=data['y'])


def torch_training_data(design, train):
    import torch
    arrays=design.matrix(train)
    require((arrays['player']>=0).all() and (arrays['pc']>=0).all() and (arrays['champion']>=0).all(), 'Training identities must be known.')
    require(train.row_kind.eq('played').all() and train.split.eq('train').all(), 'Alternative/held-out labels are forbidden.')
    tensors={k:torch.as_tensor(v,dtype=torch.long if k in ('group','player','pc','champion') else torch.float64) for k,v in arrays.items()}
    tensors.update(champion_basis=torch.tensor(design.champion_basis,dtype=torch.float64),
                   group_basis=torch.tensor(design.group_basis,dtype=torch.float64),
                   y=torch.tensor(train.win.to_numpy(dtype=float),dtype=torch.float64))
    return tensors


def fit_one(train, design, out, config, *, grouped):
    """Sequential, resumable NUTS chains. Never optimize or approximate implicitly."""
    try:
        import torch
        import pyro
        from pyro.infer import MCMC,NUTS
    except ImportError as error:
        raise RuntimeError('Fitting requires pyro-ppl==1.9.2 in this notebook environment. No package was installed and no approximation was substituted.') from error
    require(pyro.__version__ == '1.9.2', f'Exact pyro-ppl==1.9.2 required; found {pyro.__version__}.')
    previous_threads=torch.get_num_threads()
    from hlr_model.tradeoff_helpers import rank_mcmc_diagnostics
    settings=config.sampling
    model='grouped' if grouped else 'shared'
    out=Path(out)/'fits'/model;out.mkdir(parents=True,exist_ok=True)
    manifest=dict(schema=SCHEMA,model=model,inference='Pyro NUTS',backend_version=pyro.__version__,
        torch_version=torch.__version__,torch_threads=config.torch_threads,
        sampler_source_sha256=sha(inspect.getsource(fit_one).encode()),
        sampling=settings,seed=config.seed,target_accept=config.target_accept,
        training_sha256=sha(train.to_json(orient='split',index=False).encode()),
        design=design.metadata(),model_source_sha256=sha(inspect.getsource(pyro_model).encode()))
    manifest=json.loads(canonical(manifest))
    file=out/'fit_manifest.json'
    if file.is_file():require(json.loads(file.read_text())==manifest,'Fit cache differs from this exact training/model configuration.')
    else:atomic_json(file,manifest)
    data=torch_training_data(design,train)
    dimensions=dict(players=len(design.players),pcs=len(design.pcs),champions=len(design.champions))
    chain_samples=[];infos=[]
    torch.set_num_threads(config.torch_threads)
    try:
        for chain in range(settings['chains']):
            print(f'{model}: chain {chain+1}/{settings["chains"]}', flush=True)
            sample_file=out/f'chain_{chain}.npz';info_file=out/f'chain_{chain}.json'
            if sample_file.is_file() and info_file.is_file():
                info=json.loads(info_file.read_text())
                require(sha(sample_file.read_bytes())==info['samples_sha256'],'Cached chain checksum changed.')
            else:
                pyro.set_rng_seed(config.seed+chain+(1000 if grouped else 0))
                small=['alpha','beta','beta_rank','beta_role','beta_exp','log_k',
                       'log_sigma_player','log_sigma_pc','log_sigma_champion']
                if grouped:small+=['log_sigma_group','group_free']
                kernel=NUTS(pyro_model,target_accept_prob=config.target_accept,
                            max_tree_depth=settings['max_tree_depth'],full_mass=[tuple(small)],
                            # 25 smoke warmup steps cannot estimate the dense mass block reliably.
                            # Retain full mass adaptation for the 600-step full profile.
                            adapt_mass_matrix=config.profile=='full',jit_compile=False)
                sampler=MCMC(kernel,num_samples=settings['draws'],warmup_steps=settings['warmup'],num_chains=1,disable_progbar=False)
                tick=time.perf_counter();sampler.run(data,dimensions,grouped=grouped)
                values={k:v.detach().cpu().numpy() for k,v in sampler.get_samples().items()}
                buffer=io.BytesIO();np.savez_compressed(buffer,**values);atomic_bytes(sample_file,buffer.getvalue())
                diagnostics=sampler.diagnostics()
                info=dict(chain=chain,seconds=time.perf_counter()-tick,
                          divergences=len(diagnostics['divergences']['chain 0']),
                          acceptance_rate=float(diagnostics['acceptance rate']['chain 0']),
                          samples_sha256=sha(sample_file.read_bytes()))
                atomic_json(info_file,info)
            with np.load(sample_file) as archive:
                values={k:archive[k] for k in archive.files}
            require(values and all(v.shape[0] == settings['draws'] and np.isfinite(v).all()
                                   for v in values.values()), 'Invalid retained chain shapes or nonfinite draws.')
            chain_samples.append(values)
            infos.append(info)
            print(f'{model}: chain {chain+1} saved/verified; divergences={info["divergences"]}, '
                  f'acceptance={info["acceptance_rate"]:.3f}', flush=True)
    finally:
        torch.set_num_threads(previous_threads)
    samples={k:np.stack([s[k] for s in chain_samples]) for k in chain_samples[0]}
    diagnostic=rank_mcmc_diagnostics(samples)
    finite=all(np.isfinite(v).all() for v in samples.values())
    passed=bool(finite and np.isfinite(diagnostic[['rank_split_Rhat','bulk_ESS','tail_ESS']]).all().all()
        and diagnostic.rank_split_Rhat.max()<config.max_rhat
        and diagnostic.bulk_ESS.min()>config.min_ess and diagnostic.tail_ESS.min()>config.min_ess
        and sum(i['divergences'] for i in infos)==0 and min(i['acceptance_rate'] for i in infos)>=.8)
    status='smoke_only' if config.profile=='smoke' else 'converged' if passed else 'unconverged'
    diagnostic.to_csv(out/'sampling_diagnostics.csv',index=False)
    atomic_json(out/'fit_status.json',dict(status=status,diagnostic_gate_passed=passed,chains=infos,
        thresholds=dict(max_rhat=config.max_rhat,min_bulk_tail_ess=config.min_ess,max_divergences=0,min_acceptance=.8),
        limitations='Tree-depth/BFMI are not emitted by this wrapper; source benchmarks are treated as fixed. Smoke fits never authorize primary conclusions.'))
    return dict(model=model,design=design,samples=samples,status=status,diagnostics=diagnostic,chain_info=infos)


def flatten_draws(result):
    return {k:v.reshape((-1,)+v.shape[2:]) for k,v in result['samples'].items()}


def predict_draws(design, frame, samples, *, mode='personalized', seed=2026):
    """Conditional general vs personalized predictions, with identity-stable new latents.

    ``samples`` has draw as its first axis (already flattened across chains).
    No labels, published gain, clipped group score, or adjusted percentage enters eta.
    """
    require(mode in ('general','personalized'),'Specify the prediction quantity explicitly.')
    arrays=design.matrix(frame);D=len(samples['alpha'])
    def normal(kind,key):
        hashed=int.from_bytes(hashlib.sha256(repr((seed,kind,key)).encode()).digest()[:8],'little')
        return np.random.default_rng(hashed).normal(size=D)
    sigma_c=np.exp(samples['log_sigma_champion'])
    champ=(samples['champion_free']@design.champion_basis.T)*sigma_c[:,None] if len(design.champions)>1 else np.zeros((D,1))
    baseline=np.column_stack([champ[:,idx] if idx>=0 else sigma_c*normal('champion',int(cid))
                              for idx,cid in zip(arrays['champion'],frame.model_champion_id)])
    slope=np.repeat(samples['beta_exp'][:,None],len(GROUPS),axis=1)
    if 'group_free' in samples:
        slope=slope+np.exp(samples['log_sigma_group'])[:,None]*(samples['group_free']@design.group_basis.T)
    eta=(samples['alpha'][:,None]+samples['beta']@arrays['X'].T+samples['beta_rank']@arrays['rank'].T
         +samples['beta_role']@arrays['role'].T+baseline
         +slope[:,arrays['group']]*arrays['n'][None,:]/(arrays['n'][None,:]+np.exp(samples['log_k'])[:,None]))
    if mode=='personalized':
        player=np.column_stack([samples['player_z'][:,idx] if idx>=0 else normal('player',p)
                                for idx,p in zip(arrays['player'],frame.puuid)])
        pc=np.column_stack([samples['pc_z'][:,idx] if idx>=0 else normal('pc',(p,int(c)))
                            for idx,p,c in zip(arrays['pc'],frame.puuid,frame.model_champion_id)])
        eta+=np.exp(samples['log_sigma_player'])[:,None]*player+np.exp(samples['log_sigma_pc'])[:,None]*pc
    return expit(eta)


def posterior_mean(result, frame, mode):
    draws=flatten_draws(result)
    return np.concatenate([predict_draws(result['design'],frame.iloc[i:i+128],draws,mode=mode).mean(axis=0)
                           for i in range(0,len(frame),128)]) if len(frame) else np.array([])


def evaluate(frame, probability, model, mode, split):
    require(len(frame)==len(probability) and frame.row_kind.eq('played').all(),'Only observed held-out outcomes can be scored.')
    probability=np.clip(np.asarray(probability),1e-12,1-1e-12);y=frame.win.to_numpy()
    score=dict(model=model,mode=mode,split=split,rows=len(frame),players=frame.puuid.nunique(),matches=frame.match_id.nunique(),
               log_loss=log_loss(y,probability,labels=[0,1]),brier=brier_score_loss(y,probability),
               roc_auc=roc_auc_score(y,probability) if len(np.unique(y))>1 else np.nan)
    by_player=pd.DataFrame(dict(puuid=frame.puuid.to_numpy(),
        loss=-(y*np.log(probability)+(1-y)*np.log1p(-probability)),brier=(y-probability)**2)).groupby('puuid').mean()
    score.update(equal_player_log_loss=by_player.loss.mean(),equal_player_brier=by_player.brier.mean())
    calibration=frame[KEY].copy();calibration['y']=y;calibration['p']=probability
    calibration['bin']=pd.cut(probability,np.linspace(0,1,11),include_lowest=True)
    calibration=calibration.groupby('bin',observed=True).agg(rows=('y','size'),players=('puuid','nunique'),matches=('match_id','nunique'),predicted=('p','mean'),observed=('y','mean')).reset_index()
    calibration['bin']=calibration['bin'].astype(str)
    calibration['model']=model;calibration['mode']=mode;calibration['split']=split
    return score,calibration


def comparison_pairs(prepared):
    """Existing all-role main, restricted to actual same-role eligibility; no new requests."""
    p=prepared['played'];a=prepared['alternatives']
    targets=p.loc[p.split.eq('test') & p.usable_data]
    columns=KEY+['target_row_id','timestamp','lane','model_rank','model_champion_id','mastery_group',
                 'n_recorded','historical_wr_beta11','log_total','matchup_probability']
    played=p[columns+['same_role_prior_games','usable_data']].copy()
    alts=a[columns+['prior_same_role_games','usable_data']].rename(columns={'prior_same_role_games':'same_role_prior_games'})
    pool=pd.concat([played,alts],ignore_index=True)
    require(not pool.duplicated(KEY+['model_champion_id']).any(),'Prediction choice pool repeats a champion.')
    indexed={k:f for k,f in pool.groupby(KEY,sort=False)}
    left=[];right=[];records=[]
    for target in targets.itertuples(index=False):
        key=(target.puuid,target.match_id);options=indexed[key]
        for reference,cid in [('played',target.champion_id),('main',target.prior_main_champion_id)]:
            reference_rows=options.loc[options.model_champion_id.eq(cid)&options.usable_data]
            if reference=='main':reference_rows=reference_rows.loc[reference_rows.same_role_prior_games.gt(0)]
            if len(reference_rows)!=1:continue
            ref=reference_rows.iloc[0]
            choices=options.loc[options.usable_data & options.same_role_prior_games.gt(0)
                & options.model_champion_id.ne(cid) & options.matchup_probability.gt(ref.matchup_probability)]
            if reference=='played':choices=choices.loc[choices.model_champion_id.ne(target.champion_id)]
            for _,candidate in choices.iterrows():
                left.append(ref[columns].to_dict());right.append(candidate[columns].to_dict())
                records.append(dict(puuid=target.puuid,match_id=target.match_id,target_row_id=target.target_row_id,
                    reference=reference,reference_champion_id=int(cid),alternative_champion_id=int(candidate.model_champion_id),
                    reference_experience=ref.n_recorded,alternative_experience=candidate.n_recorded,
                    reference_group=ref.mastery_group,alternative_group=candidate.mastery_group,
                    reference_historical_wr=ref.historical_wr_beta11,alternative_historical_wr=candidate.historical_wr_beta11,
                    reference_raw_matchup_probability=ref.matchup_probability,
                    alternative_raw_matchup_probability=candidate.matchup_probability,
                    raw_matchup_advantage_pp=100*(candidate.matchup_probability-ref.matchup_probability)))
    return pd.DataFrame(records),pd.DataFrame(left,columns=columns),pd.DataFrame(right,columns=columns)


def recommendations(prepared, results):
    pairs,reference,alternative=comparison_pairs(prepared)
    train=prepared['played'].loc[prepared['played'].split.eq('train')&prepared['played'].usable_data]
    pc_counts=train.groupby(['puuid','model_champion_id']).size()
    champion_counts=train.groupby('model_champion_id').size()
    group_players=train.groupby('mastery_group').puuid.nunique()
    output=[]
    for name,result in results.items():
        draws=flatten_draws(result)
        for mode in ('general','personalized'):
            for i in range(0,len(pairs),128):
                r=reference.iloc[i:i+128];a=alternative.iloc[i:i+128]
                pr=predict_draws(result['design'],r,draws,mode=mode)
                pa=predict_draws(result['design'],a,draws,mode=mode)
                delta=100*(pa-pr);low,high=np.quantile(delta,[.025,.975],axis=0)
                batch=pairs.iloc[i:i+128].copy()
                batch['model']=name;batch['mode']=mode
                batch['reference_probability']=pr.mean(axis=0);batch['alternative_probability']=pa.mean(axis=0)
                batch['difference_pp']=delta.mean(axis=0);batch['low_pp']=low;batch['high_pp']=high
                batch['probability_alternative_higher']=(delta>0).mean(axis=0)
                batch['preference']=np.select([low>0,high<0],['Alternative','Reference'],default='Uncertain')
                batch['point_prefers_alternative']=batch.difference_pp.gt(0)
                batch['reference_pc_training_supported']=[(p,int(c)) in result['design'].pci for p,c in zip(r.puuid,r.model_champion_id)]
                batch['alternative_pc_training_supported']=[(p,int(c)) in result['design'].pci for p,c in zip(a.puuid,a.model_champion_id)]
                for prefix,choice in [('reference',r),('alternative',a)]:
                    batch[prefix+'_pc_training_games']=[int(pc_counts.get((p,int(c)),0)) for p,c in zip(choice.puuid,choice.model_champion_id)]
                    batch[prefix+'_champion_training_games']=[int(champion_counts.get(int(c),0)) for c in choice.model_champion_id]
                    batch[prefix+'_group_training_players']=[int(group_players.get(g,0)) for g in choice.mastery_group]
                    batch[prefix+'_sparse_training_support']=(batch[prefix+'_pc_training_games'].lt(5)
                        |batch[prefix+'_champion_training_games'].lt(20)|batch[prefix+'_group_training_players'].lt(5))
                output.append(batch)
    if not output:return pd.DataFrame(),pd.DataFrame()
    frame=pd.concat(output,ignore_index=True);summaries=[]
    for key,f in frame.groupby(['model','mode','reference']):
        events={'point_prefers_alternative':f.point_prefers_alternative,'interval_favors_alternative':f.preference.eq('Alternative'),'uncertain':f.preference.eq('Uncertain')}
        for event,values in events.items():
            weighted=f[KEY].copy();weighted['value']=values.astype(float).to_numpy()
            target=weighted.groupby(KEY).value.mean()
            summaries.append(dict(model=key[0],mode=key[1],reference=key[2],event=event,pairs=len(f),
                targets=len(f[KEY].drop_duplicates()),players=f.puuid.nunique(),
                pair_weighted=weighted.value.mean(),equal_target=target.mean(),equal_player=target.groupby(level='puuid').mean().mean()))
    return frame,pd.DataFrame(summaries)


def curve_outputs(train, results, out):
    import matplotlib.pyplot as plt
    all_curves=[]
    for model,result in results.items():
        samples=flatten_draws(result);design=result['design'];D=len(samples['alpha'])
        reference=train[list(CONTINUOUS)].median().to_numpy(dtype=float)
        x=(reference-design.mean)/design.scale
        baseline=samples['alpha']+samples['beta']@x  # Bronze/top reference, zero champion/player/PC baseline.
        slopes=np.repeat(samples['beta_exp'][:,None],len(GROUPS),axis=1)
        if 'group_free' in samples:
            slopes+=np.exp(samples['log_sigma_group'])[:,None]*(samples['group_free']@design.group_basis.T)
        fig,axes=plt.subplots(3,4,figsize=(15,10),squeeze=False)
        for j,group in enumerate(GROUPS):
            ax=axes.flat[j];support=train.loc[train.mastery_group.eq(group)]
            if support.empty:
                ax.set(title=group+' — no training support');ax.axis('off');continue
            lower=float(support.n_recorded.min())
            # A reference slice cannot have more champion games than total games.
            upper=min(float(support.n_recorded.quantile(.95)),float(np.expm1(reference[1])))
            if upper<lower:
                ax.set(title=group+' — no support at the common total-games context');ax.axis('off');continue
            grid=np.linspace(lower,upper,60) if upper>lower else np.array([lower])
            probability=expit(baseline[:,None]+slopes[:,j,None]*grid[None,:]/(grid[None,:]+np.exp(samples['log_k'])[:,None]))
            low,mean,high=np.quantile(probability,.025,axis=0),probability.mean(axis=0),np.quantile(probability,.975,axis=0)
            ax.plot(grid,mean);ax.fill_between(grid,low,high,alpha=.2)
            rug=support.loc[support.n_recorded.between(lower,upper),'n_recorded']
            ax.plot(rug,np.full(len(rug),.02),'|',transform=ax.get_xaxis_transform(),alpha=.15)
            ax.set(title=f'{group}: {len(support)} games / {support.puuid.nunique()} players',xlabel='Recorded prior all-role games',ylabel='Win probability',ylim=(0,1))
            for k,n in enumerate(grid):all_curves.append(dict(model=model,group=group,n_recorded=n,mean=mean[k],low=low[k],high=high[k],
                training_games=len(support),training_players=support.puuid.nunique(),support_min=lower,support_95th=upper,
                context='Bronze/top; training-median raw matchup/log_total/historical WR; zero champion/player/PC baselines'))
        axes.flat[-1].axis('off');fig.suptitle(model+' — common reference context; ranges stop at group training 95th percentile')
        fig.tight_layout();fig.savefig(Path(out)/'figures'/f'{model}_group_curves.png',dpi=140);plt.close(fig)
    frame=pd.DataFrame(all_curves);frame.to_parquet(Path(out)/'group_experience_curves.parquet',index=False)
    return frame


def run_modeling(prepared, study_metadata, out, config, *, enabled=False):
    if not enabled:
        print('RUN_MODELING=False: preparation/EDA only; no optimizer, sampler, collection, or recommendations were run.')
        return None
    require(study_metadata.get('modeling_ready') is True,'External study/grouping input is unresolved. Preparation is saved; fitting is blocked.')
    train=prepared['played'].loc[prepared['played'].split.eq('train')&prepared['played'].usable_data].copy()
    limit=config.sampling['train_match_limit']
    if limit is not None:
        matches=train.sort_values(['timestamp','match_id']).match_id.drop_duplicates().head(limit)
        train=train.loc[train.match_id.isin(matches)].copy()
    design=Design(train)
    results={}
    for grouped in (False,True):
        name='grouped' if grouped else 'shared'
        print(f'{name}: {len(train)} observed training rows; NUTS settings {config.sampling}')
        results[name]=fit_one(train,design,out,config,grouped=grouped)
        if config.profile=='full' and results[name]['status']!='converged':
            raise RuntimeError(f'{name} fit did not pass sampling diagnostics. Chains/diagnostics were saved; no evaluation or recommendations authorized.')
    if config.profile=='smoke':
        print('SMOKE ONLY: same likelihood/inference method on a small training subset. No primary performance or recommendation conclusions are produced.')
        return results
    freeze=Path(out)/'evaluation_freeze.json'
    if not freeze.exists():atomic_json(freeze,dict(created_at=stamp(),training_rows=len(train),
        models=['shared','grouped'],preprocessing='identical training-only design',refit_on_heldout=False,
        inference='NUTS; both models passed convergence gate',validation_start=config.validation_start,test_start=config.test_start))
    validation=evaluate_split(prepared,results,out,split='validation')
    return dict(results=results,**validation)


def evaluate_split(prepared, results, out, *, split):
    """Score one frozen held-out period; verified completed outputs are reused."""
    require(split in ('validation','test'), 'Choose one held-out period.')
    require(set(results)=={'shared','grouped'} and all(r['status']=='converged' for r in results.values()),
            'Both full models must pass diagnostics before scoring.')
    out=Path(out)
    require(split!='test' or (out/'validation_decision.json').is_file(),
            'Record the validation comparison and model decision before opening test outcomes.')
    files=[f'{split}_metrics.csv',f'{split}_calibration.csv',f'{split}_predictions.parquet']
    heldout=prepared['played'].loc[prepared['played'].split.eq(split)&prepared['played'].usable_data]
    require(len(heldout)>0,'No usable held-out rows; do not claim a successful evaluation.')
    identity=dict(rows_sha256=sha(heldout.to_json(orient='split',index=False).encode()),
        helper_sha256=sha(Path(__file__).read_bytes()),
        posterior={name:{'design':r['design'].metadata(),'samples':{
            k:sha(np.ascontiguousarray(v).tobytes()) for k,v in r['samples'].items()}}
            for name,r in results.items()})
    identity=json.loads(canonical(identity))
    manifest=out/f'{split}_evaluation.json'
    if manifest.is_file():
        saved=json.loads(manifest.read_text())
        require(saved['identity']==identity,'Held-out evaluation inputs or frozen fit changed.')
        require(all(sha((out/n).read_bytes())==h for n,h in saved['output_sha256'].items()),
                'Held-out output checksum mismatch.')
        return dict(metrics=pd.read_csv(out/files[0]),calibration=pd.read_csv(out/files[1]),
                    predictions=pd.read_parquet(out/files[2]))
    scores=[];calibrations=[];predictions=[]
    for name,result in results.items():
        for mode in ('general','personalized'):
            probability=posterior_mean(result,heldout.drop(columns='win'),mode)
            score,calibration=evaluate(heldout,probability,name,mode,split)
            scores.append(score);calibrations.append(calibration)
            f=heldout[KEY+['target_row_id']].copy();f['probability']=probability
            f['model']=name;f['mode']=mode;f['split']=split;predictions.append(f)
    metrics=pd.DataFrame(scores);calibration=pd.concat(calibrations,ignore_index=True)
    predictions=pd.concat(predictions,ignore_index=True)
    metrics.to_csv(out/files[0],index=False);calibration.to_csv(out/files[1],index=False)
    predictions.to_parquet(out/files[2],index=False)
    atomic_json(manifest,dict(created_at=stamp(),identity=identity,
        output_sha256={n:sha((out/n).read_bytes()) for n in files}))
    return dict(metrics=metrics,calibration=calibration,predictions=predictions)


def finalize_modeling(prepared, model_results, out, config, *, decision):
    """Freeze a reviewed validation decision, then score final test and report curves/pairs."""
    require(config.profile=='full' and bool(decision.strip()),'A full fit and written validation decision are required.')
    out=Path(out);results=model_results['results']
    validation=evaluate_split(prepared,results,out,split='validation')
    record=dict(decision=decision,validation_sha256=sha((out/'validation_evaluation.json').read_bytes()),
                config=asdict(config),models=['shared','grouped'],selection_uses_test=False)
    record=json.loads(canonical(record));file=out/'validation_decision.json'
    if file.exists():require(json.loads(file.read_text())['record']==record,'The pre-test decision is already frozen.')
    else:atomic_json(file,dict(created_at=stamp(),record=record))
    test=evaluate_split(prepared,results,out,split='test')
    metrics=pd.concat([validation['metrics'],test['metrics']],ignore_index=True)
    calibration=pd.concat([validation['calibration'],test['calibration']],ignore_index=True)
    predictions=pd.concat([validation['predictions'],test['predictions']],ignore_index=True)
    metrics.to_csv(Path(out)/'evaluation_metrics.csv',index=False)
    calibration.to_csv(Path(out)/'calibration.csv',index=False)
    predictions.to_parquet(Path(out)/'heldout_predictions.parquet',index=False)
    pairs,summary=recommendations(prepared,results)
    pairs.to_parquet(Path(out)/'paired_recommendations.parquet',index=False)
    summary.to_csv(Path(out)/'recommendation_summary.csv',index=False)
    train=prepared['played'].loc[prepared['played'].split.eq('train')&prepared['played'].usable_data]
    curve_outputs(train,results,out)
    import matplotlib.pyplot as plt
    fig,ax=plt.subplots(figsize=(6,5))
    for key,f in calibration.loc[calibration.split.eq('test')].groupby(['model','mode']):
        ax.plot(f.predicted,f.observed,'o-',label=' / '.join(key))
    ax.plot([0,1],[0,1],'--',color='gray');ax.set(xlabel='Mean predicted probability',ylabel='Observed win fraction',title='Held-out calibration — bin counts in calibration.csv')
    ax.legend(fontsize=8);fig.tight_layout();fig.savefig(Path(out)/'figures/test_calibration.png',dpi=140);plt.close(fig)
    atomic_json(Path(out)/'modeling_complete.json',dict(status='both_models_converged_and_evaluated',created_at=stamp(),
        primary_evaluation='chronological held-out test; retrospective benchmark-assisted',limits=LIMITS))
    return dict(results=results,metrics=metrics,calibration=calibration,recommendation_summary=summary)
