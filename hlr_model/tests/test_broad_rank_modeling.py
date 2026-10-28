"""Offline-only preparation and inference-layout checks; never launch a sampler."""
import json
from dataclasses import replace
from pathlib import Path
from unittest.mock import Mock

import numpy as np
import pandas as pd
import pytest
import requests
from openpyxl import Workbook

from hlr_model import broad_rank_modeling as m


@pytest.fixture(autouse=True)
def forbid_network(monkeypatch):
    monkeypatch.setattr(requests.sessions.Session,'request',lambda *a,**kw:pytest.fail('Network forbidden'))


def study_book(path, *, wrong_units=False, unknown=False):
    book=Workbook();intro=book.active;intro.title='Data Intro'
    intro['C1']='Fixture, not the real study';intro['C2']='https://jackjgaming.substack.com/p/fixture'
    intro['C2'].hyperlink=intro['C2'].value
    book.create_sheet('Easiest to Learn');sheet=book.create_sheet('Best to Master')
    headers=['Most Common Lane','Champion Name','Medium Mastery Win Rate','High Mastery Win Rate','High Mastery Ratio']
    for col,value in enumerate(headers,3):sheet.cell(1,col,value)
    for row,(name,medium,high) in enumerate([('Annie',.40,.50),('Ahri',.45,.55),('Fizz',.51,.50),('Milio',.52,'low data'),('Nunu & Willump',.50,.50)],2):
        for col,value in enumerate(['Mid',name,medium,high,None],3):sheet.cell(row,col,value)
        sheet.cell(row,5).number_format='0.00%';sheet.cell(row,6).number_format='0.00%'
    if wrong_units:sheet.cell(2,6,50)
    if unknown:sheet.cell(2,4,'Possibly Some Champion')
    book.save(path)


def test_workbook_units_signed_gain_clipping_low_data_and_explicit_alias(tmp_path):
    file=tmp_path/'fixture.xlsx';study_book(file)
    cfg=replace(m.Config(),study_file=str(file),study_sha256=m.sha(file.read_bytes()))
    lookup,rows,meta=m.load_groups(tmp_path,cfg)
    f=lookup.set_index('champion_id')
    assert f.loc[1,'raw_gain_pp']==10 and f.loc[103,'raw_gain_pp']==10
    assert f.loc[1,'mastery_group']==f.loc[103,'mastery_group']  # Exact tied gains stay together.
    assert f.loc[105,'raw_gain_pp']==-1 and f.loc[105,'grouping_score']==0 and f.loc[105,'mastery_group']=='J'
    assert f.loc[902,'mastery_group']=='Unclassified' and np.isnan(f.loc[902,'grouping_score'])
    assert f.loc[902,'classification_reason']=='low_data'
    assert f.loc[20,'mastery_group']=='J'
    assert meta['modeling_ready'] and not meta['external_gain_in_prediction']
    assert rows.study_high_original.str.contains('low data').sum()==1


@pytest.mark.parametrize('wrong_units,unknown',[(True,False),(False,True)])
def test_invalid_units_or_ambiguous_names_are_not_guessed(tmp_path,wrong_units,unknown):
    file=tmp_path/'bad.xlsx';study_book(file,wrong_units=wrong_units,unknown=unknown)
    cfg=replace(m.Config(),study_file=str(file),study_sha256=m.sha(file.read_bytes()))
    with pytest.raises(ValueError):m.load_groups(tmp_path,cfg)


def test_absent_workbook_is_unclassified_without_imputation(tmp_path):
    lookup,rows,meta=m.load_groups(tmp_path,replace(m.Config(),study_file='absent.xlsx'))
    assert lookup.mastery_group.eq('Unclassified').all() and lookup.grouping_score.isna().all()
    assert rows.empty and not meta['modeling_ready']


def test_group_boundaries_are_deterministic_tie_preserving_and_can_be_empty():
    values=[9,9,8,7,6,5,4,3,2,1,0,-1]
    f=pd.DataFrame(dict(champion_id=range(len(values)),study_usable=True,grouping_score=np.maximum(0,values)))
    first=m.assign_groups(f);second=m.assign_groups(f.sample(frac=1,random_state=8)).sort_index()
    pd.testing.assert_series_equal(first.mastery_group,second.mastery_group)
    assert first.loc[0,'mastery_group']==first.loc[1,'mastery_group']=='A'
    assert first.loc[9,'mastery_group']=='I' and first.loc[10:,'mastery_group'].eq('J').all()
    equal=pd.DataFrame(dict(champion_id=range(20),study_usable=True,grouping_score=1.))
    assert m.assign_groups(equal).mastery_group.nunique()==1
    assert len(set('ABCDEFGHI')-set(m.assign_groups(equal).mastery_group))==8


def model_frame():
    return pd.DataFrame(dict(puuid=['p','p','q','q'],match_id=['a','b','c','d'],target_row_id=['a','b','c','d'],
        timestamp=[100,200,300,400],
        row_kind='played',split='train',usable_data=True,win=[1,0,1,0],model_rank=['BRONZE','SILVER','BRONZE','GOLD'],
        lane=['mid','top','mid','adc'],model_champion_id=[1,103,1,105],mastery_group=['A','J','A','Unclassified'],
        n_recorded=[0,2,4,10],matchup_probability=[.4,.5,.6,.55],log_total=[1,2,3,4],historical_wr_beta11=[.5,.6,.5,.3]))


def fake_draws(design, *, grouped=False, count=200):
    rng=np.random.default_rng(71)
    values=dict(alpha=np.full(count,.1),beta=np.tile([.4,.2,.1],(count,1)),
        beta_rank=np.zeros((count,len(m.RANKS)-1)),beta_role=np.zeros((count,len(m.ROLES)-1)),
        beta_exp=np.full(count,.8),log_k=np.full(count,np.log(30)),
        log_sigma_player=np.full(count,np.log(.3)),log_sigma_pc=np.full(count,np.log(.15)),
        log_sigma_champion=np.full(count,np.log(.3)),player_z=np.zeros((count,len(design.players))),
        pc_z=np.zeros((count,len(design.pcs))),champion_free=np.zeros((count,len(design.champions)-1)))
    if grouped:
        values.update(log_sigma_group=np.full(count,np.log(.3)),group_free=rng.normal(size=(count,len(m.GROUPS)-1)))
    return values


def test_training_only_design_and_alternatives_never_supply_labels():
    train=model_frame();design=m.Design(train)
    heldout=train.copy();heldout['split']='test';heldout['matchup_probability']=1
    before=design.mean.copy();design.matrix(heldout)
    np.testing.assert_array_equal(design.mean,before)
    with pytest.raises(ValueError):m.Design(heldout)
    bad=train.copy();bad['row_kind']='alternative'
    with pytest.raises(ValueError):m.Design(bad)


def test_published_gain_is_not_a_predictor_and_group_j_can_have_an_effect():
    frame=model_frame();design=m.Design(frame);draws=fake_draws(design)
    changed=frame.assign(raw_gain_pp=9999,grouping_score=9999)
    np.testing.assert_allclose(m.predict_draws(design,frame,draws),m.predict_draws(design,changed,draws))
    zero=frame.iloc[[1]].copy();zero['n_recorded']=0
    experienced=zero.copy();experienced['n_recorded']=50
    assert zero.mastery_group.iloc[0]=='J'
    assert (m.predict_draws(design,experienced,draws)>m.predict_draws(design,zero,draws)).all()


def test_general_zero_effects_and_unseen_personalized_effects_include_pooled_uncertainty():
    frame=model_frame();design=m.Design(frame);draws=fake_draws(design)
    unseen=frame.iloc[[0]].copy();unseen['puuid']='new-player';unseen['model_champion_id']=99999
    general=m.predict_draws(design,unseen,draws,mode='general')
    personal=m.predict_draws(design,unseen,draws,mode='personalized')
    assert personal.std()>general.std() and personal.std()>0
    repeated=pd.concat([unseen,unseen],ignore_index=True)
    prediction=m.predict_draws(design,repeated,draws,mode='personalized')
    np.testing.assert_array_equal(prediction[:,0],prediction[:,1])
    draws['player_z'][:]=2;draws['pc_z'][:]=1
    assert not np.allclose(m.predict_draws(design,frame,draws,mode='personalized'),m.predict_draws(design,frame,draws,mode='general'))


def test_same_match_stays_together_and_unfinished_outcomes_are_purged():
    cfg=replace(m.Config(),validation_start='2026-01-02T00:00:00+00:00',test_start='2026-01-03T00:00:00+00:00')
    c1=pd.Timestamp(cfg.validation_start).value//10**6;c2=pd.Timestamp(cfg.test_start).value//10**6
    f=pd.DataFrame(dict(puuid=['p','q','p','q','p'],match_id=['shared','shared','early','v','t'],target_row_id=range(5),
        timestamp=[c1-10,c1-10,c1-100,c1+1,c2+1],end_timestamp=[c1-1,c1+2,c1-20,c2+1,c2+20]))
    split=m.split_assignments(f,cfg)
    assert split.loc[split.match_id.eq('shared'),'split'].eq('purged_train_boundary').all()
    assert split.loc[split.match_id.eq('early'),'split'].eq('train').all()
    assert split.loc[split.match_id.eq('v'),'split'].eq('purged_validation_boundary').all()
    assert split.groupby('match_id').split.nunique().eq(1).all()


def test_metric_counts_calibration_and_single_class_auc():
    frame=model_frame();score,cal=m.evaluate(frame,[.8,.4,.7,.2],'shared','personalized','test')
    assert score['rows']==4 and score['players']==2 and score['matches']==4
    assert score['log_loss']>0 and score['brier']>0 and score['roc_auc']==1
    assert cal.rows.sum()==4
    one=frame.loc[frame.win.eq(1)];score,_=m.evaluate(one,[.8,.7],'shared','general','test')
    assert np.isnan(score['roc_auc'])


def test_default_switch_never_calls_fitter_or_imports_backend(tmp_path,monkeypatch):
    fitter=Mock(side_effect=AssertionError('Fitting forbidden'));monkeypatch.setattr(m,'fit_one',fitter)
    assert m.run_modeling({}, {}, tmp_path,m.Config(),enabled=False) is None
    fitter.assert_not_called()
    with pytest.raises(ValueError):m.run_modeling({},dict(modeling_ready=False),tmp_path,m.Config(),enabled=True)
    fitter.assert_not_called()


def test_unconverged_fit_cannot_authorize_evaluation(tmp_path,monkeypatch):
    train=model_frame()
    fitter=Mock(return_value=dict(status='unconverged'));monkeypatch.setattr(m,'fit_one',fitter)
    scorer=Mock(side_effect=AssertionError('Evaluation forbidden'));monkeypatch.setattr(m,'evaluate',scorer)
    with pytest.raises(RuntimeError,match='did not pass sampling diagnostics'):
        m.run_modeling(dict(played=train),dict(modeling_ready=True),tmp_path,m.Config(),enabled=True)
    assert fitter.call_count==1
    scorer.assert_not_called()


def test_smoke_is_small_nuts_and_never_presents_primary_conclusions(tmp_path,monkeypatch):
    cfg=replace(m.Config(),profile='smoke')
    assert cfg.sampling==dict(chains=2,warmup=25,draws=25,max_tree_depth=4,train_match_limit=24)
    fitter=Mock(return_value=dict(status='smoke_only'));monkeypatch.setattr(m,'fit_one',fitter)
    scorer=Mock(side_effect=AssertionError('Smoke evaluation forbidden'));monkeypatch.setattr(m,'evaluate',scorer)
    result=m.run_modeling(dict(played=model_frame()),dict(modeling_ready=True),tmp_path,cfg,enabled=True)
    assert set(result)=={'shared','grouped'} and fitter.call_count==2
    scorer.assert_not_called()


def test_real_workbook_and_full_preparation_preserve_rows_and_chronology():
    root=Path(__file__).resolve().parents[2];cfg=m.Config()
    bundle=m.load_source(root,cfg);groups,study,meta=m.load_groups(root,cfg)
    prepared=m.prepare(bundle,groups,cfg)
    assert meta['available'] and meta['modeling_ready']
    assert study.study_usable.sum()==study.study_high_probability.notna().sum()
    assert study.loc[study.classification_reason.eq('low_data'),'mastery_group'].eq('Unclassified').all()
    assert np.allclose(study.loc[study.study_usable,'raw_gain_pp'],100*(study.loc[study.study_usable,'study_high_probability']-study.loc[study.study_usable,'study_medium_probability']))
    for original,output in [('played_target_games','played'),('eligible_alternatives','alternatives')]:
        assert len(prepared[output])==len(bundle['frames'][original])
        assert prepared[output].target_row_id.equals(bundle['frames'][original].target_row_id)
    p,a=prepared['played'],prepared['alternatives']
    assert p.groupby('match_id').split.nunique().eq(1).all()
    assert a.latest_all_role_prior_end_timestamp.lt(a.timestamp).all()
    assert p.prior_latest_end_timestamp.dropna().lt(p.loc[p.prior_latest_end_timestamp.notna(),'timestamp']).all()
    assert 'win' not in a and a.candidate_outcome.isna().all()
    assert p.loc[p.benchmark_sparse & p.benchmark_valid_rate,'usable_data'].all()
    assert p.loc[~p.benchmark_valid_rate,'matchup_probability'].isna().all()


@pytest.mark.parametrize('grouped',[False,True])
def test_pyro_density_and_numpy_prediction_parameterizations_match_without_sampling(grouped,monkeypatch):
    pyro=pytest.importorskip('pyro');torch=pytest.importorskip('torch')
    from pyro import poutine
    from pyro.infer import MCMC
    monkeypatch.setattr(MCMC,'run',lambda *a,**kw:pytest.fail('Sampler forbidden even in this backend check'))
    train=model_frame();design=m.Design(train);draws=fake_draws(design,grouped=grouped,count=1)
    values={k:torch.tensor(v[0],dtype=torch.float64) for k,v in draws.items()}
    data=m.torch_training_data(design,train)
    dimensions=dict(players=len(design.players),pcs=len(design.pcs),champions=len(design.champions))
    trace=poutine.trace(poutine.condition(m.pyro_model,data=values)).get_trace(data,dimensions,grouped=grouped)
    logits=trace.nodes['outcomes']['fn'].base_dist.logits
    np.testing.assert_allclose(torch.sigmoid(logits).numpy(),m.predict_draws(design,train,draws,mode='personalized')[0],atol=1e-12)
    manual=torch.distributions.Bernoulli(logits=logits).log_prob(data['y']).sum()
    priors={k:(0.,1.) for k in values}
    for k in ('alpha','beta'):priors[k]=(0.,2.)
    priors['log_k']=(np.log(30),1.)
    for k in ('log_sigma_player','log_sigma_champion','log_sigma_group'):
        if k in values:priors[k]=(np.log(.3),.7)
    priors['log_sigma_pc']=(np.log(.15),.5)
    for name,value in values.items():
        loc,scale=priors[name];manual+=torch.distributions.Normal(loc,scale).log_prob(value).sum()
    assert abs(trace.log_prob_sum().item()-manual.item())<1e-5


def test_played_and_existing_all_role_main_are_distinct_reference_scopes():
    played=model_frame().iloc[[0]].copy()
    played['split']='test';played['same_role_prior_games']=8;played['prior_main_champion_id']=103
    played['champion_id']=1;played['n_recorded']=20;played['matchup_probability']=.5
    alts=pd.concat([played,played],ignore_index=True).drop(columns=['win','same_role_prior_games','champion_id','prior_main_champion_id'])
    alts['row_kind']='alternative';alts['model_champion_id']=[103,105];alts['mastery_group']=['J','Unclassified']
    alts['n_recorded']=[40,2];alts['prior_same_role_games']=[12,1];alts['matchup_probability']=[.45,.6]
    pairs,left,right=m.comparison_pairs(dict(played=played,alternatives=alts))
    assert pairs.groupby('reference').size().to_dict()=={'main':2,'played':1}
    assert pairs.raw_matchup_advantage_pp.gt(0).all()
    assert pairs.loc[pairs.reference.eq('main'),'reference_champion_id'].eq(103).all()
    assert set(pairs.loc[pairs.reference.eq('main'),'alternative_champion_id'])=={1,105}
    assert 'win' not in left and 'win' not in right


def test_matched_postfit_outputs_with_mocked_draws_use_no_fitter_or_sampler(tmp_path,monkeypatch):
    """Exercise future reporting on synthetic posteriors, never real fitted results."""
    frames=[]
    for split in ('train','validation','test'):
        f=model_frame();f['split']=split;f['match_id']=split+'-'+f.match_id;f['target_row_id']=f.match_id
        f['champion_id']=f.model_champion_id;f['prior_main_champion_id']=f.model_champion_id
        f['same_role_prior_games']=f.n_recorded;frames.append(f)
    played=pd.concat(frames,ignore_index=True)
    alt=played.loc[played.match_id.eq('test-b')].drop(columns=['win','champion_id','prior_main_champion_id','same_role_prior_games']).copy()
    alt['row_kind']='alternative';alt['model_champion_id']=1;alt['mastery_group']='A'
    alt['n_recorded']=1;alt['prior_same_role_games']=1;alt['matchup_probability']=.8
    def fake_fit(train,design,out,config,*,grouped):
        draws=fake_draws(design,grouped=grouped,count=20)
        return dict(model='grouped' if grouped else 'shared',design=design,status='converged',
                    samples={k:v[None,...] for k,v in draws.items()})
    monkeypatch.setattr(m,'fit_one',fake_fit)
    (tmp_path/'figures').mkdir()
    result=m.run_modeling(dict(played=played,alternatives=alt),dict(modeling_ready=True),tmp_path,m.Config(),enabled=True)
    assert len(result['metrics'])==4 and result['metrics'].split.eq('validation').all()
    assert not (tmp_path/'test_evaluation.json').exists()
    assert not (tmp_path/'paired_recommendations.parquet').exists()
    with pytest.raises(ValueError,match='Record the validation'):
        m.evaluate_split(dict(played=played),result['results'],tmp_path,split='test')
    result=m.finalize_modeling(dict(played=played,alternatives=alt),result,tmp_path,m.Config(),
        decision='Synthetic validation fixture: retain both fixed specifications for final comparison.')
    assert len(result['metrics'])==8  # 2 models × 2 prediction quantities × 2 held-out periods.
    assert result['metrics'].rows.eq(4).all()
    pairs=pd.read_parquet(tmp_path/'paired_recommendations.parquet')
    assert 'win' not in pairs and pairs.raw_matchup_advantage_pp.gt(0).all()
    assert set(pairs['mode'])=={'general','personalized'} and set(pairs['reference'])=={'played','main'}
    curves=pd.read_parquet(tmp_path/'group_experience_curves.parquet')
    assert curves.n_recorded.le(curves.support_95th).all()
    assert curves.n_recorded.le(np.expm1(frames[0].log_total.median())).all()
    assert set(curves['group'])=={'A','J','Unclassified'}
    assert (tmp_path/'figures/test_calibration.png').is_file()
    assert (tmp_path/'modeling_complete.json').is_file()
    monkeypatch.setattr(m,'posterior_mean',lambda *a,**kw:pytest.fail('Completed test must reuse verified predictions'))
    m.evaluate_split(dict(played=played),result['results'],tmp_path,split='test')
    with (tmp_path/'test_metrics.csv').open('a') as f:f.write('\ntampered')
    with pytest.raises(ValueError,match='checksum mismatch'):
        m.evaluate_split(dict(played=played),result['results'],tmp_path,split='test')
